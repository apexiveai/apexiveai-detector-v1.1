from pathlib import Path
from functools import lru_cache
from collections import Counter
import csv, hashlib, json, re, traceback, zipfile, io, math, shutil, subprocess, tempfile, sys, time
import os
import pymupdf
import cv2
import numpy as np
from PIL import Image
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from skimage.metrics import structural_similarity as ssim
import xml.etree.ElementTree as ET

MEDIA_EXT={'.png','.jpg','.jpeg','.bmp','.webp','.tif','.tiff','.gif','.wmf','.emf','.svg'}
APP_RE=re.compile(r'(?<![A-Z0-9])(?:[A-Z]{1,8}/)?[T7I]\s*/\s*(\d{4})\s*/\s*(\d{3,10})(?!\d)',re.I)
APP_LABEL_RE=re.compile(r'\b(?:APPLICATION|APP)\s*(?:NO\.?|NUMBER)\s*[:#-]?\s*([A-Z0-9][A-Z0-9/-]{2,31})\b',re.I)
MARK_WORDS={'MARK','LOGO','DEVICE','TRADEMARK','TRADEMARKS','LABEL','BRAND','FIGURE'}
PDF_FIELD_CODE_RE=re.compile(r'\(\d{3}\)')
MARK_HEADERS={word.lower() for word in MARK_WORDS}|{'reproductionofmark','mainobjective','mainobjectives','emainobjective','emainobjectives'}
STANDALONE_IMAGE_EXTENSIONS={'.png','.jpg','.jpeg','.bmp','.webp','.tif','.tiff','.gif'}
MAX_ASSETS_PER_SOURCE=int(os.getenv('MAX_ASSETS_PER_SOURCE','6000'))
MAX_WORKSHEET_ROWS=int(os.getenv('MAX_WORKSHEET_ROWS','250000'))
MAX_WORKSHEET_COLUMNS=int(os.getenv('MAX_WORKSHEET_COLUMNS','256'))
MIN_VISUAL_SIMILARITY_SCORE=float(os.getenv('MIN_VISUAL_SIMILARITY_SCORE','0.70'))
if not 0 <= MIN_VISUAL_SIMILARITY_SCORE <= 1:
    raise ValueError('MIN_VISUAL_SIMILARITY_SCORE must be between 0 and 1.')

def norm_app(v):
    if v is None:return None
    m=APP_RE.search(str(v).upper())
    return f'T/{m.group(1)}/{m.group(2)}' if m else None

def application_text(value):
    if value is None:return ''
    if isinstance(value,float) and value.is_integer():value=int(value)
    return str(value).strip()

def application_key(value):
    text=application_text(value)
    if not text:return ''
    return norm_app(text) or re.sub(r'[^A-Z0-9]+','',text.upper())

def application_number_status(asset_a,asset_b):
    raw_a=application_text(asset_a.get('application_raw') or asset_a.get('application_no'))
    raw_b=application_text(asset_b.get('application_raw') or asset_b.get('application_no'))
    if not raw_a or not raw_b:return 'missing'
    if raw_a==raw_b:return 'exact'
    if application_key(raw_a)==application_key(raw_b):return 'normalized'
    return 'mismatch'

def application_pair_counts(assets_a,assets_b):
    raw_a=Counter(
        application_text(asset.get('application_raw') or asset.get('application_no'))
        for asset in assets_a
        if application_key(asset.get('application_raw') or asset.get('application_no'))
    )
    raw_b=Counter(
        application_text(asset.get('application_raw') or asset.get('application_no'))
        for asset in assets_b
        if application_key(asset.get('application_raw') or asset.get('application_no'))
    )
    by_key_a={}
    by_key_b={}
    for raw,count in raw_a.items():
        by_key_a.setdefault(application_key(raw),Counter())[raw]=count
    for raw,count in raw_b.items():
        by_key_b.setdefault(application_key(raw),Counter())[raw]=count
    exact=0
    normalized=0
    for key in by_key_a.keys() & by_key_b.keys():
        left=by_key_a[key]
        right=by_key_b[key]
        same=sum(left[raw]*right[raw] for raw in left.keys() & right.keys())
        exact+=same
        normalized+=sum(left.values())*sum(right.values())-same
    valid_pairs=sum(raw_a.values())*sum(raw_b.values())
    total_pairs=len(assets_a)*len(assets_b)
    return {
        'exact':exact,
        'normalized':normalized,
        'mismatch':valid_pairs-exact-normalized,
        'missing':total_pairs-valid_pairs,
        'pairs_checked':total_pairs,
    }

def sha_file(p):
    h=hashlib.sha256()
    with open(p,'rb') as f:
        for b in iter(lambda:f.read(1048576),b''):h.update(b)
    return h.hexdigest()

def cv(data):
    try:
        if isinstance(data,(bytes,bytearray)):return cv2.imdecode(np.frombuffer(data,np.uint8),cv2.IMREAD_COLOR)
        return cv2.imdecode(np.fromfile(str(data),np.uint8),cv2.IMREAD_COLOR)
    except Exception:return None

def render_windows_metafile(src,dst):
    if sys.platform!='win32':return None
    import ctypes
    from ctypes import wintypes

    class StartupInput(ctypes.Structure):
        _fields_=[
            ('GdiplusVersion',wintypes.UINT),
            ('DebugEventCallback',ctypes.c_void_p),
            ('SuppressBackgroundThread',wintypes.BOOL),
            ('SuppressExternalCodecs',wintypes.BOOL),
        ]

    gdiplus=ctypes.WinDLL('gdiplus')
    token=ctypes.c_size_t()
    startup=StartupInput(1,None,False,False)
    if gdiplus.GdiplusStartup(ctypes.byref(token),ctypes.byref(startup),None)!=0:return None
    metafile=ctypes.c_void_p()
    bitmap=ctypes.c_void_p()
    graphics=ctypes.c_void_p()
    try:
        if gdiplus.GdipCreateMetafileFromFile(str(src),ctypes.byref(metafile))!=0:return None
        width=wintypes.UINT();height=wintypes.UINT()
        if gdiplus.GdipGetImageWidth(metafile,ctypes.byref(width))!=0:return None
        if gdiplus.GdipGetImageHeight(metafile,ctypes.byref(height))!=0:return None
        if not width.value or not height.value:return None
        if width.value>12000 or height.value>12000:return None
        if gdiplus.GdipCreateBitmapFromScan0(width.value,height.value,0,0x26200A,None,ctypes.byref(bitmap))!=0:return None
        if gdiplus.GdipGetImageGraphicsContext(bitmap,ctypes.byref(graphics))!=0:return None
        if gdiplus.GdipGraphicsClear(graphics,0xFFFFFFFF)!=0:return None
        if gdiplus.GdipDrawImageRectI(graphics,metafile,0,0,width.value,height.value)!=0:return None
        png_encoder=ctypes.c_byte*16
        encoder_id=png_encoder(0x06,0xF4,0x7C,0x55,0x04,0x1A,0xD3,0x11,0x9A,0x73,0x00,0x00,0xF8,0x1E,0xF3,0x2E)
        if gdiplus.GdipSaveImageToFile(bitmap,str(dst),ctypes.byref(encoder_id),None)!=0:return None
        return cv(dst)
    finally:
        if graphics:gdiplus.GdipDeleteGraphics(graphics)
        if bitmap:gdiplus.GdipDisposeImage(bitmap)
        if metafile:gdiplus.GdipDisposeImage(metafile)
        gdiplus.GdiplusShutdown(token)

def render_vector_bytes(data,ext):
    suffix=ext if ext.startswith('.') else '.'+ext
    if suffix=='.svg':
        data=sanitize_svg(data)
    with tempfile.TemporaryDirectory() as td:
        src=Path(td)/('asset'+suffix);dst=Path(td)/'asset.png';src.write_bytes(data)
        if suffix in {'.emf','.wmf'}:
            im=render_windows_metafile(src,dst)
            if im is not None:return im
        rsvg=shutil.which('rsvg-convert');magick=shutil.which('magick');ink=shutil.which('inkscape')
        cmds=[]
        if rsvg and suffix=='.svg':
            cmds.append([rsvg,'--format=png','--output='+str(dst),'--width=2048','--height=2048','--keep-aspect-ratio',str(src)])
        if magick:cmds.append([magick,'-density','300',str(src),'-background','white','-alpha','remove','-alpha','off',str(dst)])
        if ink and suffix in {'.svg','.emf','.wmf'}:cmds.append([ink,str(src),'--export-type=png','--export-filename='+str(dst)])
        for cmd in cmds:
            try:
                p=subprocess.run(cmd,capture_output=True,timeout=45)
                if p.returncode==0 and dst.exists():
                    im=cv(dst)
                    if im is not None:return im
            except (OSError,subprocess.TimeoutExpired):
                continue
    return None

def sanitize_svg(data):
    if len(data)>25*1024*1024:
        raise ValueError('SVG file exceeds the configured 25 MB limit.')
    upper=data.upper()
    if b'<!DOCTYPE' in upper or b'<!ENTITY' in upper:
        raise ValueError('SVG DTDs and entity declarations are not supported.')
    try:
        root=ET.fromstring(data)
    except ET.ParseError as exc:
        raise ValueError('The uploaded SVG is damaged or invalid.') from exc
    if root.tag.rsplit('}',1)[-1].lower()!='svg':
        raise ValueError('The uploaded file is not an SVG document.')
    for element in root.iter():
        tag=element.tag.rsplit('}',1)[-1].lower()
        if tag in {'script','foreignobject'}:
            raise ValueError('SVG scripts and embedded foreign content are not supported.')
        for name,value in element.attrib.items():
            attribute=name.rsplit('}',1)[-1].lower()
            text=value.strip()
            if attribute.startswith('on') or attribute=='base':
                raise ValueError('SVG event handlers and external base URLs are not supported.')
            if attribute in {'href','src'} and text and not text.startswith('#'):
                if not re.fullmatch(r'data:image/(?:png|jpe?g|gif|webp);base64,[A-Za-z0-9+/=]+',text,re.I):
                    raise ValueError('SVG external image or resource references are not supported.')
            if '@import' in text.lower():
                raise ValueError('SVG external stylesheets are not supported.')
            for reference in re.findall(r'url\(\s*[\'"]?([^\'")]+)',text,re.I):
                if not reference.strip().startswith('#'):
                    raise ValueError('SVG external image or resource references are not supported.')
    return ET.tostring(root,encoding='utf-8',xml_declaration=True)

def svg_text(data):
    root=ET.fromstring(sanitize_svg(data))
    values=list(root.itertext())
    values.extend(
        value for element in root.iter()
        for key,value in element.attrib.items()
        if key.rsplit('}',1)[-1].lower() in {'aria-label','title','description'}
    )
    return ' '.join(values)

def gray(img):return cv2.cvtColor(img,cv2.COLOR_BGR2GRAY) if len(img.shape)==3 else img.copy()

def logo_mask(g):
    g=cv2.normalize(g,None,0,255,cv2.NORM_MINMAX);g=cv2.GaussianBlur(g,(3,3),0)
    _,m=cv2.threshold(g,0,255,cv2.THRESH_BINARY_INV+cv2.THRESH_OTSU)
    if np.mean(g)<127:m=cv2.bitwise_not(m)
    m=cv2.morphologyEx(m,cv2.MORPH_OPEN,cv2.getStructuringElement(cv2.MORPH_ELLIPSE,(3,3)))
    return m

def canonical(img,size=256):
    if img is None:return None
    g=gray(img);m=logo_mask(g)
    n,lab,stats,_=cv2.connectedComponentsWithStats(m,8);clean=np.zeros_like(m);min_area=max(6,int(m.size*.00003))
    for i in range(1,n):
        if stats[i,cv2.CC_STAT_AREA]>=min_area:clean[lab==i]=255
    ys,xs=np.where(clean>0)
    if len(xs)>20:
        pad=max(4,int(min(g.shape)*.03));g=g[max(0,ys.min()-pad):min(g.shape[0],ys.max()+pad+1),max(0,xs.min()-pad):min(g.shape[1],xs.max()+pad+1)]
    g=cv2.resize(g,(size,size),interpolation=cv2.INTER_AREA)
    if np.mean(g)<127:g=255-g
    return g

def canonical_color(img,size=256):
    if img is None:return None
    source=img if len(img.shape)==3 else cv2.cvtColor(img,cv2.COLOR_GRAY2BGR)
    g=gray(source);mask=logo_mask(g);ys,xs=np.where(mask>0)
    if len(xs)>20:
        pad=max(4,int(min(g.shape)*.03))
        source=source[
            max(0,ys.min()-pad):min(g.shape[0],ys.max()+pad+1),
            max(0,xs.min()-pad):min(g.shape[1],xs.max()+pad+1),
        ]
    return cv2.resize(source,(size,size),interpolation=cv2.INTER_AREA)

def color_signature(img):
    if img is None:return np.zeros(24,dtype=np.float32)
    source=img if len(img.shape)==3 else cv2.cvtColor(img,cv2.COLOR_GRAY2BGR)
    mask=logo_mask(gray(source))
    signature=[]
    for channel in cv2.split(source):
        hist=cv2.calcHist([channel],[0],mask,[8],[0,256]).flatten().astype(np.float32)
        total=float(hist.sum())
        signature.extend(hist/total if total else np.zeros(8,dtype=np.float32))
    return np.asarray(signature,dtype=np.float32)

def layout_signature(img):
    g=canonical(img)
    if g is None:return np.zeros(16,dtype=np.float32)
    mask=logo_mask(g)
    return cv2.resize(mask,(4,4),interpolation=cv2.INTER_AREA).astype(np.float32).flatten()/255

def bits(a):
    v=0
    for x in a.flatten().astype(np.uint8):v=(v<<1)|int(x)
    return f'{v:0{(a.size+3)//4}x}'

def hamming(a,b):return (int(a,16)^int(b,16)).bit_count()

def hashes(g):
    sm=cv2.resize(g,(32,32),interpolation=cv2.INTER_AREA).astype(np.float32);d=cv2.dct(sm)[:8,:8];med=np.median(d.flatten()[1:])
    ph=bits((d>med).astype(np.uint8));dhg=cv2.resize(g,(9,8),interpolation=cv2.INTER_AREA);dh=bits((dhg[:,1:]>dhg[:,:-1]).astype(np.uint8));ahg=cv2.resize(g,(8,8),interpolation=cv2.INTER_AREA);ah=bits((ahg>ahg.mean()).astype(np.uint8))
    return ph,dh,ah

def fingerprint_from_canonical(g):
    m=logo_mask(g)
    ph,dh,ah=hashes(g)
    return {'sha256':hashlib.sha256(g.tobytes()).hexdigest(),'mask_sha256':hashlib.sha256(m.tobytes()).hexdigest(),'phash':ph,'dhash':dh,'ahash':ah}

def fp(img):
    g=canonical(img)
    if g is None:return {}
    fingerprint=fingerprint_from_canonical(g)
    color=canonical_color(img)
    fingerprint.update({
        'color_sha256':hashlib.sha256(color.tobytes()).hexdigest(),
        'color_signature':color_signature(img),
        'layout_signature':layout_signature(img),
        'contour':contour_signature(img),
    })
    return fingerprint

def fast_score(a,b):
    ga,gb=canonical(a),canonical(b)
    if ga is None or gb is None:return 0
    fa,fb=hashes(ga),hashes(gb)
    return .55*(1-hamming(fa[0],fb[0])/64)+.30*(1-hamming(fa[1],fb[1])/64)+.15*(1-hamming(fa[2],fb[2])/64)

def contour_signature(img):
    g=gray(img);m=logo_mask(g);contours,_=cv2.findContours(m,cv2.RETR_EXTERNAL,cv2.CHAIN_APPROX_SIMPLE)
    if not contours:return None
    contour=max(contours,key=cv2.contourArea);hu=cv2.HuMoments(cv2.moments(contour)).flatten();hu=np.sign(hu)*np.log10(np.abs(hu)+1e-30);x,y,w,h=cv2.boundingRect(contour)
    return hu,cv2.contourArea(contour)/(g.shape[0]*g.shape[1]),w/g.shape[1],h/g.shape[0]

def contour_similarity_from_signatures(x,y):
    if not x or not y:return 0
    hu=.0
    for i in range(7):hu+=min(abs(float(x[0][i]-y[0][i]))/10,1)
    hu=1-hu/7;ar=1-abs(x[1]-y[1])/max(x[1],y[1],1e-6);sh=1-(abs(x[2]-y[2])+abs(x[3]-y[3]))/2
    return max(0,min(1,.55*hu+.25*ar+.20*sh))

def contour_similarity(a,b):
    return contour_similarity_from_signatures(contour_signature(a),contour_signature(b))

def prepare_verification(img):
    g=canonical(img)
    if g is None:return {}
    sift=cv2.SIFT_create(nfeatures=600);keypoints,descriptors=sift.detectAndCompute(g,None)
    fingerprint=fp(img)
    return {
        'canonical':g,
        'fingerprint':fingerprint,
        'mask':logo_mask(g),
        'keypoints':keypoints,
        'descriptors':descriptors,
        'contour':fingerprint['contour'],
        'color_signature':fingerprint['color_signature'],
        'layout_signature':fingerprint['layout_signature'],
    }

def contour_similarities(signature,signatures):
    if signature is None or not len(signatures):
        return np.zeros(len(signatures),dtype=np.float32)
    hu=1-np.minimum(np.abs(signatures[:,:7]-signature[0])/10,1).mean(axis=1)
    area=1-np.abs(signatures[:,7]-signature[1])/np.maximum(np.maximum(signatures[:,7],signature[1]),1e-6)
    shape=1-(np.abs(signatures[:,8]-signature[2])+np.abs(signatures[:,9]-signature[3]))/2
    return np.clip(.55*hu+.25*area+.20*shape,0,1)

def fingerprint_upper_bounds(fingerprint,hashes_array,colors_array,layouts_array,contours_array):
    if not len(hashes_array):return np.zeros(0,dtype=np.float32)
    query_hashes=np.asarray([
        int(fingerprint.get('phash','0'),16),
        int(fingerprint.get('dhash','0'),16),
        int(fingerprint.get('ahash','0'),16),
    ],dtype=np.uint64)
    similarities=1-np.bitwise_count(np.bitwise_xor(hashes_array,query_hashes))/64
    ph,dh,ah=similarities.T
    query_color=fingerprint.get('color_signature',np.zeros(24,dtype=np.float32))
    color=np.sqrt(np.maximum(colors_array*query_color,0)).sum(axis=1)/3
    query_layout=fingerprint.get('layout_signature',np.zeros(16,dtype=np.float32))
    layout=1-np.abs(layouts_array-query_layout).mean(axis=1)
    contour=contour_similarities(fingerprint.get('contour'),contours_array)
    return .15*ph+.07*dh+.05*ah+.15+.15*contour+.15*color+.10*layout+.08+.10

def contour_vector(signature):
    if signature is None:return np.zeros(10,dtype=np.float32)
    return np.concatenate((signature[0].astype(np.float32),np.asarray(signature[1:],dtype=np.float32)))

def application_fields(asset_a,asset_b):
    status=application_number_status(asset_a,asset_b)
    return {
        'application_a':asset_a.get('application_no',''),
        'application_b':asset_b.get('application_no',''),
        'application_raw_a':asset_a.get('application_raw',''),
        'application_raw_b':asset_b.get('application_raw',''),
        'application_number_status':status,
        'application_no_equal':status in {'exact','normalized'},
    }

def logo_match_record(job_dir,asset_a,asset_b,score,decision,origin,metrics):
    return {
        'match_id':'',
        'source_a_asset_id':asset_a['asset_id'],
        'source_b_asset_id':asset_b['asset_id'],
        'source_a_file':asset_a['source_file'],
        'source_b_file':asset_b['source_file'],
        'source_a_page':asset_a.get('page'),
        'source_b_page':asset_b.get('page'),
        'source_a_sheet':asset_a.get('sheet'),
        'source_a_row':asset_a.get('row'),
        'source_a_column':asset_a.get('column'),
        'source_b_sheet':asset_b.get('sheet'),
        'source_b_row':asset_b.get('row'),
        'source_b_column':asset_b.get('column'),
        'source_a_reference':asset_a.get('source_reference'),
        'source_b_reference':asset_b.get('source_reference'),
        'source_a_field':asset_a.get('field_label',''),
        'source_b_field':asset_b.get('field_label',''),
        **application_fields(asset_a,asset_b),
        'score':round(score,6),
        'decision':decision,
        'match_origin':origin,
        'metrics':metrics,
        'asset_a_url':f"/api/jobs/{job_dir.name}/assets/A/{asset_a['image_file']}",
        'asset_b_url':f"/api/jobs/{job_dir.name}/assets/B/{asset_b['image_file']}",
    }

def iter_match_records(path):
    with open(path,encoding='utf8') as source:
        for line in source:
            if line.strip():
                yield json.loads(line)

def screen_logo_pairs(job_dir,work,report,visual_a,visual_b,Bfp,hashes_array,colors_array,layouts_array,contours_array,update):
    matches_path=report/'matches.jsonl'
    matched_a=set()
    matched_b=set()
    decisions=Counter()
    pairs_checked=0
    pair_total=len(visual_a)*len(visual_b)
    candidate_total=0
    verified=0
    application_equal_matches=0
    last_update=time.monotonic()
    if pair_total==0:
        update(90,'No logo-image pairs were available for visual comparison.')
    with open(matches_path,'w',encoding='utf8') as match_output:
        @lru_cache(maxsize=128)
        def prepared_file(image_file):
            return prepare_verification(cv(work/image_file))

        for asset_a in visual_a:
            image_a=cv(work/asset_a['image_file'])
            fingerprint_a=fp(image_a)
            if not fingerprint_a:
                raise ValueError(f"Could not prepare extracted logo '{asset_a['image_file']}' for comparison.")
            upper_bounds=fingerprint_upper_bounds(
                fingerprint_a,hashes_array,colors_array,layouts_array,contours_array
            )
            candidate_indexes=np.flatnonzero(
                upper_bounds>=MIN_VISUAL_SIMILARITY_SCORE-1e-6
            )
            pairs_checked+=len(visual_b)
            candidate_total+=len(candidate_indexes)
            prepared_a=None
            for candidate_index in candidate_indexes:
                asset_b=Bfp[int(candidate_index)]
                fingerprint_b=asset_b['_fp']
                if fingerprint_a['color_sha256']==fingerprint_b['color_sha256']:
                    score=1.0
                    exact=True
                    detail={
                        'phash':1.0,'dhash':1.0,'ahash':1.0,'ssim':1.0,
                        'contour_similarity':1.0,'mask_iou':1.0,'mask_ssim':1.0,
                        'color_similarity':1.0,'layout_similarity':1.0,
                        'sift_good':0,'sift_inliers':0,
                        'canonical_sha256':fingerprint_a['sha256'],
                        'color_sha256':fingerprint_a['color_sha256'],
                        'mask_sha256':fingerprint_a['mask_sha256'],
                    }
                else:
                    if prepared_a is None:
                        prepared_a=prepare_verification(image_a)
                    score,detail,exact=verify_prepared(
                        prepared_a,prepared_file(asset_b['image_file'])
                    )
                if exact:
                    decision='EXACT_VISUAL_IDENTITY'
                elif score>=MIN_VISUAL_SIMILARITY_SCORE:
                    decision='VERY_HIGH_VISUAL_SIMILARITY' if score>=.82 else 'HIGH_VISUAL_SIMILARITY'
                else:
                    verified+=1
                    continue
                match=logo_match_record(
                    job_dir,asset_a,asset_b,score,decision,
                    'exact_color_fingerprint' if exact else 'exhaustive_fingerprint_screen',
                    detail,
                )
                match['match_id']=f'M-{sum(decisions.values())+1:06d}'
                match_output.write(json.dumps(match,ensure_ascii=False,separators=(',',':'))+'\n')
                matched_a.add(asset_a['asset_id'])
                matched_b.add(asset_b['asset_id'])
                decisions[decision]+=1
                application_equal_matches+=int(match['application_no_equal'])
                verified+=1
                if time.monotonic()-last_update>=3:
                    verify_progress=90+int(7*verified/max(candidate_total,1))
                    update(
                        verify_progress,
                        f'Screened {pairs_checked:,}/{pair_total:,} logo pairs; '
                        f'verified {verified:,}/{candidate_total:,} candidates; '
                        f'{sum(decisions.values()):,} qualified matches.',
                    )
                    last_update=time.monotonic()
            pair_progress=55+int(35*pairs_checked/max(pair_total,1))
            update(
                pair_progress,
                f'Screened {pairs_checked:,}/{pair_total:,} logo pairs; '
                f'verified {verified:,}/{candidate_total:,} candidates; '
                f'{sum(decisions.values()):,} qualified matches.',
            )
    return {
        'matches_path':matches_path,
        'matched_a':matched_a,
        'matched_b':matched_b,
        'pairs_checked':pairs_checked,
        'pair_total':pair_total,
        'candidate_total':candidate_total,
        'verified':verified,
        'decisions':decisions,
        'application_equal_matches':application_equal_matches,
    }

MATCH_REPORT_FIELDS=[
    'match_id','source_a_asset_id','source_b_asset_id','source_a_file','source_b_file',
    'source_a_page','source_a_sheet','source_a_row','source_a_column','source_a_reference',
    'source_b_page','source_b_sheet','source_b_row','source_b_column','source_b_reference',
    'source_a_field','source_b_field','application_a','application_raw_a',
    'application_b','application_raw_b','application_number_status','application_no_equal',
    'score','decision','match_origin',
    'metric_phash','metric_dhash','metric_ahash','metric_ssim','metric_contour_similarity',
    'metric_mask_iou','metric_mask_ssim','metric_color_similarity','metric_layout_similarity',
    'metric_sift_inliers',
]

def match_report_row(match):
    row={field:match.get(field,'') for field in MATCH_REPORT_FIELDS if not field.startswith('metric_')}
    metrics=match.get('metrics') or {}
    for field in MATCH_REPORT_FIELDS:
        if field.startswith('metric_'):
            row[field]=metrics.get(field.removeprefix('metric_'),'')
    return {
        field: (
            "'" + value
            if isinstance(value,str) and value.lstrip(' \t\r\n')[:1] in {'=','+','-','@'}
            else value
        )
        for field,value in row.items()
    }

def write_match_reports(matches_path,report):
    with open(report/'matches.csv','w',newline='',encoding='utf-8-sig') as output:
        writer=csv.DictWriter(output,fieldnames=MATCH_REPORT_FIELDS)
        writer.writeheader()
        writer.writerows(match_report_row(match) for match in iter_match_records(matches_path))
    workbook=Workbook(write_only=True)
    worksheet=workbook.create_sheet('Matches')
    worksheet.append(MATCH_REPORT_FIELDS)
    rows_in_sheet=1
    sheet_number=1
    for match in iter_match_records(matches_path):
        if rows_in_sheet>=1_048_576:
            sheet_number+=1
            worksheet=workbook.create_sheet(f'Matches {sheet_number}')
            worksheet.append(MATCH_REPORT_FIELDS)
            rows_in_sheet=1
        row=match_report_row(match)
        worksheet.append([row.get(field,'') for field in MATCH_REPORT_FIELDS])
        rows_in_sheet+=1
    workbook.save(report/'matches.xlsx')

def verify_prepared(first,second):
    if not first or not second:return 0,{},False
    ga,gb=first['canonical'],second['canonical']
    fa,fb=first['fingerprint'],second['fingerprint']
    ph=1-hamming(fa['phash'],fb['phash'])/64
    dh=1-hamming(fa['dhash'],fb['dhash'])/64
    ah=1-hamming(fa['ahash'],fb['ahash'])/64
    ss=float(ssim(ga,gb,data_range=255))
    cs=contour_similarity_from_signatures(first['contour'],second['contour'])
    color=float(np.sqrt(np.maximum(first['color_signature']*second['color_signature'],0)).sum()/3)
    layout=float(max(0,1-np.abs(first['layout_signature']-second['layout_signature']).mean()))
    k1,d1=first['keypoints'],first['descriptors'];k2,d2=second['keypoints'],second['descriptors'];good=[];inl=0
    if d1 is not None and d2 is not None:
        for pair in cv2.BFMatcher().knnMatch(d1,d2,k=2):
            if len(pair)==2 and pair[0].distance<.72*pair[1].distance:good.append(pair[0])
        if len(good)>=4:
            src=np.float32([k1[m.queryIdx].pt for m in good]).reshape(-1,1,2);dst=np.float32([k2[m.trainIdx].pt for m in good]).reshape(-1,1,2)
            _,mask=cv2.findHomography(src,dst,cv2.RANSAC,4);inl=int(mask.sum()) if mask is not None else 0
    ma,mb=first['mask'],second['mask'];best_iou=0;best_ms=0
    for ang in (0,-2,2,-4,4):
        M=cv2.getRotationMatrix2D((128,128),ang,1);rb=cv2.warpAffine(mb,M,(256,256),borderValue=0);inter=np.logical_and(ma>0,rb>0).sum();union=np.logical_or(ma>0,rb>0).sum();iou=inter/max(union,1);ms=float(ssim(ma,rb,data_range=255));
        if iou>best_iou or (iou==best_iou and ms>best_ms):best_iou,best_ms=iou,ms
    exact=fa['color_sha256']==fb['color_sha256'] or (
        best_iou>=.92 and best_ms>=.995 and ph>=.94 and ss>=.94 and color>=.98
    )
    final=(
        .15*max(ph,0)+.07*max(dh,0)+.05*max(ah,0)+.15*max(ss,0)
        +.15*cs+.15*color+.10*layout+.08*best_iou+.10*min(1,inl/12)
    )
    if exact:final=max(final,.995)
    return float(final),{
        'phash':round(ph,4),'dhash':round(dh,4),'ahash':round(ah,4),
        'ssim':round(ss,5),'contour_similarity':round(cs,5),
        'mask_iou':round(best_iou,5),'mask_ssim':round(best_ms,5),
        'color_similarity':round(color,5),'layout_similarity':round(layout,5),
        'sift_good':len(good),'sift_inliers':inl,
        'canonical_sha256':fa['sha256'],'color_sha256':fa['color_sha256'],
        'mask_sha256':fa['mask_sha256'],
    },exact

def verify(a,b):
    return verify_prepared(prepare_verification(a),prepare_verification(b))

def app_metadata(text):
    m=APP_RE.search(text or '')
    return f'T/{m.group(1)}/{m.group(2)}' if m else ''

def app_raw_metadata(text):
    match=APP_RE.search(text or '')
    if match:return match.group(0)
    labeled=APP_LABEL_RE.search(text or '')
    return labeled.group(1) if labeled else ''

def page_applications(page):
    applications=[]
    for block in page.get_text('blocks'):
        text=block[4] if len(block)>4 else ''
        for match in APP_RE.finditer(text or ''):
            applications.append((
                (block[1]+block[3])/2,
                f'T/{match.group(1)}/{match.group(2)}',
                match.group(0),
            ))
    return applications

def page_text_words(page):
    try:return page.get_text('words')
    except:return []

def ocr_page_fields(page):
    try:
        import pytesseract
        from pytesseract import Output
        scale=1.5
        pix=page.get_pixmap(matrix=pymupdf.Matrix(scale,scale),alpha=False)
        image=Image.open(io.BytesIO(pix.tobytes('png')))
        data=pytesseract.image_to_data(image,output_type=Output.DICT,config='--psm 6')
    except pytesseract.TesseractNotFoundError:
        raise RuntimeError('Scanned PDF text needs the Tesseract OCR executable, which is not available.')
    text_words=[];codes=[];lines=[]
    for index,raw in enumerate(data.get('text',[])):
        token=str(raw or '').strip()
        if not token:continue
        x0=float(data['left'][index])/scale
        y0=float(data['top'][index])/scale
        x1=x0+float(data['width'][index])/scale
        y1=y0+float(data['height'][index])/scale
        text_words.append((x0,y0,x1,y1,token))
        lines.append(token)
        if re.fullmatch(r'[\(\[]?(\d{3})[\)\]]?',token) and x0<page.rect.width*.35:
            code=token.strip('()[]')
            codes.append((pymupdf.Rect(x0,y0,x1,y1),code))
    return [rect for rect,code in codes if code=='540'],text_words,' '.join(lines)

def normalized_mark_text(value):
    return re.sub(r'\s+',' ',str(value or '')).strip().casefold()

def normalize_asset_record(asset,path,side,index):
    asset.setdefault('asset_id',f'{side}-{index:06d}')
    asset.setdefault('source_file',path.name)
    asset.setdefault('source_type',path.suffix.lower().lstrip('.'))
    asset.setdefault('page',None)
    asset.setdefault('sheet',None)
    asset.setdefault('row',None)
    asset.setdefault('column',None)
    asset.setdefault('application_no','')
    asset.setdefault('application_raw',application_text(asset.get('application_no')))
    asset.setdefault('company','')
    asset.setdefault('mark','')
    asset.setdefault('mark_text','')
    asset.setdefault('mark_content_type','image' if asset.get('image_file') else 'text')
    asset.setdefault('field_code','')
    asset.setdefault('field_label','')
    asset.setdefault('extract_method','')
    asset['source_reference']=(
        f"page {asset['page']}" if asset.get('page') is not None else
        f"{asset['sheet']}!{get_column_letter(asset['column']) if asset.get('column') else ''}{asset['row']}" if asset.get('sheet') else
        asset.get('field_label') or path.name
    )
    return asset

def image_application_metadata(img):
    try:
        import pytesseract
        rgb=cv2.cvtColor(img,cv2.COLOR_BGR2RGB) if len(img.shape)==3 else img
        text=pytesseract.image_to_string(rgb,config='--psm 6',timeout=15)
        raw=app_raw_metadata(text)
        return raw,norm_app(raw) or application_key(raw), ''
    except pytesseract.TesseractNotFoundError:
        message='Tesseract OCR is unavailable; application number could not be extracted from this image.'
        return '','',message

def standalone_image_assets(path,out,side='A'):
    try:
        with Image.open(path) as source:
            source.seek(0)
            source.load()
            im=cv2.cvtColor(np.asarray(source.convert('RGB')),cv2.COLOR_RGB2BGR)
    except (OSError,ValueError) as exc:
        raise ValueError(f"Could not decode standalone logo image '{path.name}'.") from exc
    if im is None:
        raise ValueError(f"Could not decode standalone logo image '{path.name}'.")
    application_raw,application_no,application_error=image_application_metadata(im)
    fn=f'{side}_000001.png'
    if not cv2.imwrite(str(out/fn),im):
        raise ValueError(f"Could not save normalized logo image '{path.name}'.")
    return [{
        'asset_id':f'{side}-000001',
        'source_type':'image',
        'source_file':path.name,
        'application_no':application_no,
        'application_raw':application_raw,
        'application_extraction_error':application_error,
        'mark_text':'',
        'mark_content_type':'image',
        'field_code':'',
        'field_label':'Standalone logo',
        'image_file':fn,
        'extract_method':'standalone_image_normalization',
    }]

def svg_assets(path,out,side='A'):
    data=path.read_bytes()
    safe_data=sanitize_svg(data)
    text=svg_text(data)
    application_raw=app_raw_metadata(text)
    application_error=''
    im=render_vector_bytes(safe_data,'.svg')
    if im is None:
        raise ValueError('Could not render SVG logo. Install librsvg (rsvg-convert) or a supported SVG renderer and retry.')
    if not application_raw:
        application_raw,_,application_error=image_application_metadata(im)
    application_no=norm_app(application_raw) or application_key(application_raw)
    fn=f'{side}_000001.png'
    if not cv2.imwrite(str(out/fn),im):
        raise ValueError(f"Could not save normalized SVG logo '{path.name}'.")
    return [{
        'asset_id':f'{side}-000001',
        'source_type':'svg',
        'source_file':path.name,
        'application_no':application_no,
        'application_raw':application_raw,
        'application_extraction_error':application_error,
        'mark_content_type':'image',
        'field_label':'Standalone SVG logo',
        'image_file':fn,
        'extract_method':'standalone_svg_render',
    }]

def pdf_assets(path,out,progress,side='A'):
    with pymupdf.open(str(path)) as doc:
        return _pdf_assets_from_doc(doc,path,out,progress,side)

def _pdf_assets_from_doc(doc,path,out,progress,side):
    assets=[];pages=len(doc);asset_no=0
    for pi,page in enumerate(doc):
        progress(int(pi/max(pages,1)*40),f'Extracting PDF: page {pi+1}/{pages}')
        fields=page.search_for('(540)')
        applications=page_applications(page)
        ocr_mode=False;ocr_text=''
        words=page.get_text('words')
        if not fields and not words:
            fields,words,ocr_text=ocr_page_fields(page)
            ocr_mode=bool(fields)
        if not fields:
            progress(int((pi+1)/max(pages,1)*40),f'Extracting PDF: page {pi+1}/{pages}')
            continue
        code_words=[
            word for word in words
            if (PDF_FIELD_CODE_RE.fullmatch(word[4]) or (ocr_mode and re.fullmatch(r'[\(\[]?\d{3}[\)\]]?',word[4])))
            and word[0]<page.rect.width*.35
        ]
        image_rects=[]
        for info in page.get_images(full=True):
            for rect in page.get_image_rects(info[0],transform=False):
                if rect.width<10 or rect.height<10 or rect.width*rect.height>page.rect.width*page.rect.height*.90:continue
                image_rects.append(rect)
        vector_rects=[drawing['rect'] for drawing in page.get_drawings() if drawing.get('rect')]
        for field in fields:
            if len(assets)>=MAX_ASSETS_PER_SOURCE:
                raise ValueError(f'PDF contains more than {MAX_ASSETS_PER_SOURCE} mark records.')
            next_codes=[word[1] for word in code_words if word[1]>=field.y1-2 and word[1]>field.y0]
            field_bottom=min(next_codes,default=page.rect.y1)
            field_top=max(page.rect.y0,field.y0-8)
            candidates=[
                rect for rect in image_rects
                if rect.x0>=field.x1+20 and field_top<=((rect.y0+rect.y1)/2)<field_bottom
            ]
            method='pdf_540_ocr_scan' if ocr_mode else 'pdf_540_embedded_image'
            mark_text=''
            if ocr_mode:
                bbox=pymupdf.Rect(
                    max(page.rect.x0,field.x1+8),field_top,
                    page.rect.x1,max(field_top+1,min(page.rect.y1,field_bottom)),
                )
            elif candidates:
                bbox=pymupdf.Rect(
                    min(rect.x0 for rect in candidates),min(rect.y0 for rect in candidates),
                    max(rect.x1 for rect in candidates),max(rect.y1 for rect in candidates),
                )
            else:
                candidates=[
                    rect for rect in vector_rects
                    if rect.width>2 and rect.height>2
                    and rect.width<page.rect.width*.7 and rect.height<page.rect.height*.3
                    and rect.x0>=field.x1+20 and field_top<=((rect.y0+rect.y1)/2)<field_bottom
                ]
                if candidates:
                    bbox=pymupdf.Rect(
                        min(rect.x0 for rect in candidates),min(rect.y0 for rect in candidates),
                        max(rect.x1 for rect in candidates),max(rect.y1 for rect in candidates),
                    )
                    method='pdf_540_vector_art'
                else:
                    text_words=[
                        word for word in words
                        if word[1]>=field.y1-2 and word[1]<field_bottom
                        and word[0]>=field.x0 and not PDF_FIELD_CODE_RE.fullmatch(word[4])
                    ]
                    if not text_words:continue
                    text_words.sort(key=lambda word:(round(word[1],1),word[0]))
                    mark_text=' '.join(word[4] for word in text_words).strip()
                    bbox=pymupdf.Rect(
                        min(word[0] for word in text_words),min(word[1] for word in text_words),
                        max(word[2] for word in text_words),max(word[3] for word in text_words),
                    )
                    method='pdf_540_text_mark'
            bbox=pymupdf.Rect(
                max(page.rect.x0,bbox.x0-2),max(page.rect.y0,bbox.y0-2),
                min(page.rect.x1,bbox.x1+2),min(page.rect.y1,bbox.y1+2),
            )
            pix=page.get_pixmap(matrix=pymupdf.Matrix(2,2),clip=bbox,alpha=False)
            im=cv(pix.tobytes('png'))
            if im is None:raise ValueError(f'Could not render (540) logo on PDF page {pi+1}')
            asset_no+=1;fn=f'{side}_{asset_no:06d}.png';cv2.imwrite(str(out/fn),im)
            if applications:
                nearest=min(applications,key=lambda item:abs(item[0]-(field.y0+field.y1)/2))
                application_no,application_raw=nearest[1],nearest[2]
            else:
                page_text=page.get_text('text') or ocr_text
                application_raw=app_raw_metadata(page_text)
                application_no=norm_app(application_raw) or application_key(application_raw) or app_metadata(page_text)
            assets.append({
                'asset_id':f'{side}-{asset_no:06d}','source_type':'pdf',
                'source_file':path.name,'page':pi+1,'bbox':[bbox.x0,bbox.y0,bbox.x1,bbox.y1],
                'application_no':application_no,'application_raw':application_raw,
                'field_code':'(540)',
                'field_label':'(540) text mark' if method=='pdf_540_text_mark' else '(540) logo',
                'mark_text':mark_text,
                'mark_content_type':'text' if method=='pdf_540_text_mark' else 'image',
                'image_file':fn,'extract_method':method,
            })
        progress(int((pi+1)/max(pages,1)*40),f'Extracting PDF: page {pi+1}/{pages}')
    return assets

def ooxml_images(path):
    result=[]
    with zipfile.ZipFile(path) as z:
        names=set(z.namelist());ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main';nsr='http://schemas.openxmlformats.org/officeDocument/2006/relationships';xdr='http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing';a='http://schemas.openxmlformats.org/drawingml/2006/main'
        wb=ET.fromstring(z.read('xl/workbook.xml'));rels=ET.fromstring(z.read('xl/_rels/workbook.xml.rels'));wr={x.attrib.get('Id'):x.attrib.get('Target','') for x in rels}
        for sh in wb.findall('{%s}sheets/{%s}sheet'%(ns,ns)):
            rid=sh.attrib.get('{%s}id'%nsr);target=wr.get(rid)
            if not target:continue
            sp='xl/'+target.split('xl/')[-1].lstrip('/');sprels='xl/worksheets/_rels/'+Path(sp).name+'.rels'
            if sprels not in names:continue
            sr=ET.fromstring(z.read(sp));rr=ET.fromstring(z.read(sprels));rmap={x.attrib.get('Id'):x.attrib.get('Target','') for x in rr}
            de=sr.find('.//{%s}drawing'%ns)
            if de is None:continue
            dp='xl/drawings/'+Path(rmap.get(de.attrib.get('{%s}id'%nsr),'')).name
            if dp not in names:continue
            dr=ET.fromstring(z.read(dp));drels='xl/drawings/_rels/'+Path(dp).name+'.rels';drr=ET.fromstring(z.read(drels));dmap={x.attrib.get('Id'):x.attrib.get('Target','') for x in drr}
            for anc in dr:
                frm=anc.find('{%s}from'%xdr);blip=anc.find('.//{%s}blip'%a)
                if frm is None or blip is None:continue
                row=int(frm.findtext('{%s}row'%xdr,default='0'))+1;col=int(frm.findtext('{%s}col'%xdr,default='0'))+1;rid2=blip.attrib.get('{%s}embed'%nsr);target2=dmap.get(rid2,'').replace('..','').lstrip('/');media='xl/'+target2.split('xl/')[-1].lstrip('/')
                if media in names:result.append((sh.attrib.get('name','Sheet'),row,col,z.read(media),Path(media).suffix.lower()))
    return result

def xlsx_assets(path,out,progress,side='A'):
    progress(45,'Reading workbook values and indexing embedded images')
    raw=ooxml_images(path);wb=load_workbook(path,data_only=True,read_only=True,keep_links=False);assets=[];asset_no=0
    if len(raw)>MAX_ASSETS_PER_SOURCE:
        wb.close()
        raise ValueError(f'Workbook contains more than {MAX_ASSETS_PER_SOURCE} embedded mark images.')
    image_sheets={image[0] for image in raw}
    try:
        for ws in wb.worksheets:
            max_row=ws.max_row or 0;max_col=max(ws.max_column or 0,1)
            if max_row>MAX_WORKSHEET_ROWS or max_col>MAX_WORKSHEET_COLUMNS:
                raise ValueError(
                    f"Worksheet '{ws.title}' exceeds configured row/column limits "
                    f"({MAX_WORKSHEET_ROWS} rows, {MAX_WORKSHEET_COLUMNS} columns)."
                )
            mark_col=None;mark_header_row=0;headers={}
            header_end=min(20,max_row)
            if header_end:
                for r,row_values in enumerate(ws.iter_rows(min_row=1,max_row=header_end,max_col=max_col,values_only=True),1):
                    for c,value in enumerate(row_values,1):
                        v=re.sub(r'[^a-z0-9]+','',str(value or '').lower())
                        if v in MARK_HEADERS and mark_col is None:mark_col=c;mark_header_row=r
                        if 'app' in v and ('no' in v or 'number' in v):headers.setdefault('app',c)
                        if any(k in v for k in ('company','applicant','owner')):headers.setdefault('company',c)
            if mark_col is None:
                if ws.title in image_sheets:raise ValueError(f"Worksheet '{ws.title}' contains images but has no Mark/logo/device column")
                continue
            byrow={}
            for sh,row,col,data,ext in raw:
                if sh!=ws.title or col!=mark_col or row>max_row:continue
                im=cv(data)
                if im is None and ext in {'.wmf','.emf','.svg'}: im=render_vector_bytes(data,ext)
                if im is None:
                    message=f"Could not decode embedded {ext.lstrip('.').upper()} image in worksheet '{ws.title}' at row {row}, column {col}."
                    if ext in {'.wmf','.emf','.svg'}:message+=' Install ImageMagick or Inkscape with support for this format and retry.'
                    else:message+=' The image may be damaged or use an unsupported format.'
                    raise ValueError(message)
                byrow.setdefault(row,[]).append((im,row,col,ext))
            value_columns=[mark_col,headers.get('app',1)]
            if max_col>=2:value_columns.append(headers.get('company',2))
            value_max_col=max(value_columns)
            first_row=min(mark_header_row+1,min(byrow,default=mark_header_row+1))
            values_by_row={}
            for row,row_values in enumerate(ws.iter_rows(min_row=first_row,max_row=max_row,max_col=value_max_col,values_only=True),first_row):
                mark_value=row_values[mark_col-1] if mark_col<=len(row_values) else None
                if row in byrow or (mark_value is not None and str(mark_value).strip()):
                    values_by_row[row]=row_values
            for row in sorted(values_by_row):
                items=byrow.get(row,[])
                row_values=values_by_row[row]
                cell=lambda column:row_values[column-1] if column<=len(row_values) else None
                app=application_text(cell(headers.get('app',1)))
                company=str(cell(headers.get('company',2)) or '') if max_col>=2 else ''
                mark=str(cell(mark_col) or '')
                if items:
                    for im,rr,col,ext in items:
                        if asset_no>=MAX_ASSETS_PER_SOURCE:
                            raise ValueError(f'Workbook contains more than {MAX_ASSETS_PER_SOURCE} mark records.')
                        asset_no+=1;fn=f'{side}_{asset_no:06d}.png';cv2.imwrite(str(out/fn),im);assets.append({'asset_id':f'{side}-{asset_no:06d}','source_type':'xlsx','source_file':path.name,'sheet':ws.title,'row':row,'column':col,'application_no':norm_app(app) or app,'application_raw':app,'company':company,'mark':mark,'mark_text':mark,'mark_content_type':'image','field_code':'Mark','field_label':'Mark column logo','image_file':fn,'extract_method':'xlsx_mark_column_image','media_ext':ext})
                elif mark.strip():
                    if asset_no>=MAX_ASSETS_PER_SOURCE:
                        raise ValueError(f'Workbook contains more than {MAX_ASSETS_PER_SOURCE} mark records.')
                    asset_no+=1;assets.append({'asset_id':f'{side}-{asset_no:06d}','source_type':'xlsx','source_file':path.name,'sheet':ws.title,'row':row,'column':mark_col,'application_no':norm_app(app) or app,'application_raw':app,'company':company,'mark':mark,'mark_text':mark,'mark_content_type':'text','field_code':'Mark','field_label':'Mark column text mark','extract_method':'xlsx_mark_column_text'})
    finally:
        wb.close()
    progress(50,f'Extracting XLSX: {len(assets)} marks')
    return assets

def extract(path,side,out,progress):
    ext=path.suffix.lower()
    if ext=='.pdf':
        assets=pdf_assets(path,out,progress,side)
    elif ext=='.svg':
        assets=svg_assets(path,out,side)
        progress(50,f'Normalized standalone SVG logo: {path.name}')
    elif ext in STANDALONE_IMAGE_EXTENSIONS:
        assets=standalone_image_assets(path,out,side)
        progress(50,f'Normalized standalone logo: {path.name}')
    else:
        assets=xlsx_assets(path,out,progress,side)
    return [normalize_asset_record(asset,path,side,index) for index,asset in enumerate(assets,1)]

def run(job_dir,file_a,file_b,update):
    work=job_dir/'assets';work.mkdir(parents=True,exist_ok=True);report=job_dir/'report';report.mkdir(parents=True,exist_ok=True)
    update(1,'Checking whether source files are identical')
    hash_a=sha_file(file_a);hash_b=sha_file(file_b)
    identical_sources=hash_a==hash_b
    A=extract(file_a,'A',work,update)
    visual_a=[asset for asset in A if asset.get('image_file')]
    B=extract(file_b,'B',work,update)
    visual_b=[asset for asset in B if asset.get('image_file')]
    update(52,f'Building exhaustive comparison index for {len(visual_b):,} logos')
    Bfp=[]
    for asset in visual_b:
        image=cv(work/asset['image_file'])
        fingerprint=fp(image)
        if not fingerprint:
            raise ValueError(f"Could not prepare extracted logo '{asset['image_file']}' for comparison.")
        asset['_fp']=fingerprint
        Bfp.append(asset)
    hashes_array=np.asarray([
        [int(asset['_fp'][key],16) for key in ('phash','dhash','ahash')]
        for asset in Bfp
    ],dtype=np.uint64).reshape((-1,3))
    colors_array=np.stack([asset['_fp']['color_signature'] for asset in Bfp]) if Bfp else np.zeros((0,24),dtype=np.float32)
    layouts_array=np.stack([asset['_fp']['layout_signature'] for asset in Bfp]) if Bfp else np.zeros((0,16),dtype=np.float32)
    contours_array=np.stack([contour_vector(asset['_fp']['contour']) for asset in Bfp]) if Bfp else np.zeros((0,10),dtype=np.float32)

    scan=screen_logo_pairs(
        job_dir,work,report,visual_a,visual_b,Bfp,hashes_array,
        colors_array,layouts_array,contours_array,update,
    )
    match_count=sum(scan['decisions'].values())
    update(97,f'Writing {match_count:,} exact or threshold-qualified matches and exports.')
    write_match_reports(scan['matches_path'],report)
    for a in A:a.pop('_fp',None)
    for b in B:b.pop('_fp',None)
    application_counts=application_pair_counts(A,B)
    note=(
        'Logo images are compared independently of application numbers. All possible image pairs '
        'are fingerprint-screened; only exact or threshold-qualified visual matches are stored. '
        'Application numbers are normalized and compared separately. Similarity is technical '
        'evidence for human review, not a legal infringement verdict.'
    )
    summary={
        'source_a':{'file':file_a.name,'type':file_a.suffix.lower(),'assets':len(A),'sha256':hash_a},
        'source_b':{'file':file_b.name,'type':file_b.suffix.lower(),'assets':len(B),'sha256':hash_b},
        'matched':match_count,
        'unmatched_a':len(A)-len(scan['matched_a']),
        'unmatched_b':len(B)-len(scan['matched_b']),
        'exact_identity':scan['decisions']['EXACT_VISUAL_IDENTITY'],
        'exact_text_matches':0,
        'very_high':scan['decisions']['VERY_HIGH_VISUAL_SIMILARITY'],
        'high':scan['decisions']['HIGH_VISUAL_SIMILARITY'],
        'application_number_equal_matches':scan['application_equal_matches'],
        'application_number_status_counts':application_counts,
        'application_number_used_for_matching':False,
        'pairs_checked':scan['pairs_checked'],
        'qualified_candidate_pairs_verified':scan['verified'],
        'minimum_visual_similarity_score':MIN_VISUAL_SIMILARITY_SCORE,
        'identical_sources':identical_sources,
        'analysis_skipped':False,
        'same_file_logo_scan':identical_sources,
        'engine':'Apexive AI Trademark Visual Conflict Engine v5',
        'message':note,
        'note':note,
    }
    (report/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf8')
    (report/'assets_a.json').write_text(json.dumps(A,ensure_ascii=False,indent=2),encoding='utf8')
    (report/'assets_b.json').write_text(json.dumps(B,ensure_ascii=False,indent=2),encoding='utf8')
    update(100,'Completed')
    return summary
